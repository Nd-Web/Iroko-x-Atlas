"""Workspace sync, review permissions, coverage, impacts, tasks and the export."""

from io import BytesIO

import pytest

from ingestion.models import DocumentAccess, RecordAccess
from models.compliance_graph import Control, Impact, Judgement, Link, ObligationStatus
from models.workflow import WorkflowTask
from services.compliance_graph import core, queries, review, wording
from services.compliance_graph.jobs import run_job
from services.compliance_graph.permissions import GraphNotFound, GraphPermissionError
from services.compliance_graph.visibility import for_user, for_workspace
from tests.compliance_graph import factory as f


async def extract(db, document_id, model=None):
    key, token = f.claim_job(db, "graph", document_id)
    await run_job(db, key, token, "graph", document_id, complete=model or f.FakeModel())


async def sync(db, workspace_id, model=None):
    model = model or f.FakeModel()
    key, token = f.claim_job(db, "graph_ws", workspace_id)
    await run_job(db, key, token, "graph_ws", workspace_id, complete=model)
    return model


@pytest.fixture
async def world(db):
    """The Iroko library (shared BVN circular), customer A (an MFB) and customer B (a mobile money operator)."""
    iroko = f.user(db, "iroko", role="superadmin", workspace="ws-lib")
    ada = f.user(db, "ada", role="admin", workspace="ws-a")
    ali = f.user(db, "ali", role="analyst", workspace="ws-a")
    bola = f.user(db, "bola", role="admin", workspace="ws-b")
    f.document(db, "bvn", title="BVN Enrollment for OFIs Customers", pages=[f.BVN_LETTER], workspace="ws-lib",
               shared=True, regulator="CBN", reference="OFI/DIR/CIR/GEN/17/139", source_key="regulator:CBN:bvn",
               uploaded_by="iroko")
    f.document(db, "pol", title="AML/CFT Policy", pages=[f.AML_POLICY], workspace="ws-a", role="policy",
               source_key="upload:ada:aml", uploaded_by="ada")
    f.document(db, "reg", title="AML training register Q1 2026", pages=[f.TRAINING_REGISTER], workspace="ws-a",
               role="evidence_record", uploaded_by="ada")
    f.document(db, "bpol", title="B's KYC policy", pages=[f.AML_POLICY], workspace="ws-b", role="policy",
               uploaded_by="bola")
    await extract(db, "bvn")
    await extract(db, "pol", f.FakeModel(role="policy"))
    await extract(db, "reg", f.FakeModel(role="evidence_record"))
    await extract(db, "bpol", f.FakeModel(role="policy"))
    review.set_profile(db, for_user(db, ada), ["state"])
    review.set_profile(db, for_user(db, bola), ["mmo"])
    await sync(db, "ws-a")
    await sync(db, "ws-b")
    return {"iroko": iroko, "ada": ada, "ali": ali, "bola": bola}


def scope(db, user):
    return for_user(db, user)


async def test_suggestions_follow_the_licence(db, world):
    a_links = db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses").all()
    assert a_links and all(link.basis == "suggested" and link.review_status == "proposed" for link in a_links)
    assert all(len(link.anchors) == 2 and link.rationale for link in a_links)
    evidences = db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "evidences").all()
    assert evidences and evidences[0].attributes["period_end"] == "2026-03-31"
    # A mobile money operator is not among an OFI circular's addressees: nothing is suggested to B.
    assert not db.query(Link).filter(Link.workspace_id == "ws-b").count()
    rows = queries.requirement_rows(db, scope(db, world["bola"]))
    assert {r["applicability"]["state"] for r in rows} == {"not_addressed"}


async def test_review_permissions_follow_the_split(db, world):
    link = db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses").first()
    with pytest.raises(GraphPermissionError):
        review.review_link(db, scope(db, world["ali"]), link.id, "confirm")
    with pytest.raises(GraphNotFound):
        review.review_link(db, scope(db, world["bola"]), link.id, "confirm")
    applies = db.query(Link).filter(Link.relation == "applies_to", Link.from_id == "bvn").one()
    with pytest.raises(GraphPermissionError):
        review.review_link(db, scope(db, world["ada"]), applies.id, "confirm")
    assert review.review_link(db, scope(db, world["iroko"]), applies.id, "confirm").review_status == "confirmed"
    with pytest.raises(Exception):
        review.review_link(db, scope(db, world["ada"]), link.id, "reject", note="")
    assert review.review_link(db, scope(db, world["ada"]), link.id, "confirm").review_status == "confirmed"


async def test_coverage_moves_with_confirmations(db, world):
    ada = scope(db, world["ada"])
    lineage = db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses").first().to_id

    def state():
        return next(r for r in queries.requirement_rows(db, ada) if r["lineage_id"] == lineage)["coverage"]["state"]

    assert state() in ("control_suggested", "evidence_suggested")
    for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses", Link.to_id == lineage):
        review.review_link(db, ada, link.id, "confirm")
    assert state() in ("control_confirmed_no_evidence", "evidence_suggested")
    for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "evidences"):
        review.review_link(db, ada, link.id, "confirm")
    # The register covers Q1 2026 and a linked control runs monthly: current in April, stale by October.
    from datetime import date

    applicable = {lineage: "applies"}
    assert core.coverage_of(db, ada, [lineage], applicable, today=date(2026, 4, 10))[lineage].state == "evidence_on_record"
    assert core.coverage_of(db, ada, [lineage], applicable, today=date(2026, 10, 9))[lineage].state == "evidence_out_of_date"
    review.set_requirement_status(db, ada, lineage, {"applicability": "not_applicable", "note": "We have no agents"})
    assert state() == "not_applicable"
    with pytest.raises(Exception):
        review.set_requirement_status(db, ada, lineage, {"applicability": "not_applicable", "note": ""})


async def test_missed_evidence_dates_are_marked_overdue(db, world):
    from datetime import date

    ada = scope(db, world["ada"])
    for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation.in_(["addresses", "evidences"])):
        review.review_link(db, ada, link.id, "confirm")
    db.commit()
    # Q1 2026 evidence for a monthly control: the next record was expected in early May.
    cycles = [i for i in queries.due_items(db, ada, today=date(2026, 10, 9))["items"] if i["kind"] == "control_cycle"]
    assert cycles and all(i["overdue"] for i in cycles)
    upcoming = [i for i in queries.due_items(db, ada, today=date(2026, 4, 10))["items"] if i["kind"] == "control_cycle"]
    assert upcoming and not any(i["overdue"] for i in upcoming)


async def test_manual_controls_and_links(db, world):
    ali, ada = scope(db, world["ali"]), scope(db, world["ada"])
    control = review.create_control(db, ali, {"name": "Daily sanctions screening", "frequency": "daily"})
    assert (control.basis, control.review_status) == ("manual", "proposed")
    lineage = queries.requirement_rows(db, ada)[0]["lineage_id"]
    link = review.create_link(db, ali, relation="addresses", from_type="control", from_id=control.id,
                              to_type="obligation", to_id=lineage)
    assert (link.basis, link.review_status) == ("manual", "proposed")
    again = review.create_link(db, ada, relation="addresses", from_type="control", from_id=control.id,
                               to_type="obligation", to_id=lineage)
    assert again.id == link.id and again.review_status == "confirmed"
    with pytest.raises(GraphNotFound):
        review.create_link(db, scope(db, world["bola"]), relation="addresses", from_type="control",
                           from_id=control.id, to_type="obligation", to_id=lineage)
    with pytest.raises(Exception):
        review.update_control(db, ada, control.id, {"owner_user_id": "bola"})  # not a member of ws-a


async def test_new_version_flags_confirmed_links_once_it_is_shared(db, world):
    ada = scope(db, world["ada"])
    for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses"):
        review.review_link(db, ada, link.id, "confirm")
    revised = f.BVN_LETTER.replace("monthly returns on BVN enrollment", "monthly returns on BVN enrolment")
    f.supersede(db, "bvn")
    f.document(db, "bvn2", title="BVN Enrollment for OFIs Customers", pages=[revised], workspace="ws-lib",
               shared=False, regulator="CBN", reference="OFI/DIR/CIR/GEN/17/139", source_key="regulator:CBN:bvn",
               previous_id="bvn", uploaded_by="iroko")
    await extract(db, "bvn2")
    await sync(db, "ws-a")
    # Not shared yet: the customer still sees the old version, labelled, and nothing is flagged.
    rows = queries.requirement_rows(db, ada)
    assert rows and all(r["document"]["id"] == "bvn" and r["document"]["awaiting_publication"] for r in rows)
    assert not db.query(Impact).filter(Impact.workspace_id == "ws-a").count()
    queue = queries.review_queue(db, scope(db, world["iroko"]))
    assert next(g for g in queue["groups"] if g["key"] == "share_new_version")["count"] == 1
    db.get(DocumentAccess, "bvn2").shared_regulatory = True
    db.commit()
    await sync(db, "ws-a")
    impact = db.query(Impact).filter(Impact.workspace_id == "ws-a", Impact.kind == "new_version").one()
    flagged = {item["link_id"] for item in impact.affected["flagged"]}
    assert flagged and all(db.get(Link, lid).review_status == "needs_re_review" for lid in flagged)
    task = db.query(WorkflowTask).filter(WorkflowTask.source_type == "change_impact").one()
    assert db.get(RecordAccess, ("task", task.id)).workspace_id == "ws-a"
    assert {r["document"]["id"] for r in queries.requirement_rows(db, ada)} == {"bvn2"}
    review.acknowledge_impact(db, ada, impact.id)
    assert db.get(WorkflowTask, task.id).status == "done"


async def test_rejected_amendment_releases_its_flags(db, world):
    ada, iroko = scope(db, world["ada"]), scope(db, world["iroko"])
    for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses"):
        review.review_link(db, ada, link.id, "confirm")
    f.document(db, "amend", title="BVN amendment", pages=[f.AMENDING_LETTER], workspace="ws-lib", shared=True,
               regulator="CBN", reference="OFI/DIR/CIR/GEN/19/001", uploaded_by="iroko")
    await extract(db, "amend")
    await sync(db, "ws-a")
    impact = db.query(Impact).filter(Impact.workspace_id == "ws-a", Impact.kind == "amended_by").one()
    assert impact.affected["flagged"]
    amends = db.query(Link).filter(Link.relation == "amends").one()
    review.review_link(db, iroko, amends.id, "reject", note="This letter only cites the circular")
    await sync(db, "ws-a")
    db.refresh(impact)
    assert impact.status == "withdrawn"
    assert all(db.get(Link, item["link_id"]).review_status == "confirmed" for item in impact.affected["flagged"])


async def test_a_change_older_than_the_confirmation_flags_nothing(db, world):
    # The amendment is in the library before the team reviews its mapping.
    f.document(db, "amend", title="BVN amendment", pages=[f.AMENDING_LETTER], workspace="ws-lib", shared=True,
               regulator="CBN", reference="OFI/DIR/CIR/GEN/19/001", uploaded_by="iroko")
    await extract(db, "amend")
    await sync(db, "ws-a")
    ada = scope(db, world["ada"])
    confirmed = [review.review_link(db, ada, link.id, "confirm")
                 for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.relation == "addresses")]
    db.commit()
    await sync(db, "ws-a")
    assert confirmed and all(db.get(Link, link.id).review_status == "confirmed" for link in confirmed)
    assert all(not impact.affected["flagged"] for impact in db.query(Impact).filter(Impact.workspace_id == "ws-a"))


async def test_the_judge_knows_which_document_each_text_is_in(db, world):
    db.query(Judgement).delete()
    db.commit()
    model = await sync(db, "ws-a")
    asked = [c["payload"] for c in model.calls if "judgements" in c["properties"]]
    assert asked
    control_side = next(p["control"] for p in asked if "control" in p)
    assert control_side["context"].startswith("AML/CFT Policy")  # "this policy" is that document
    assert all(t["context"].startswith("BVN") for p in asked if "control" in p for t in p["targets"])


async def test_unchanged_pairs_are_never_judged_twice(db, world):
    judgements = db.query(Judgement).count()
    model = await sync(db, "ws-a")
    assert not [c for c in model.calls if "judgements" in c["properties"]]
    assert db.query(Judgement).count() == judgements


async def test_retired_controls_and_archived_sources_become_impacts(db, world):
    ada = scope(db, world["ada"])
    control = db.query(Control).filter(Control.workspace_id == "ws-a").first()
    for link in db.query(Link).filter(Link.workspace_id == "ws-a", Link.from_id == control.id):
        review.review_link(db, ada, link.id, "confirm")
    review.update_control(db, ada, control.id, {"status": "retired"})
    from models.database import Document

    db.get(Document, "reg").status = "archived"
    db.commit()
    await sync(db, "ws-a")
    kinds = {i.kind for i in db.query(Impact).filter(Impact.workspace_id == "ws-a")}
    assert {"control_retired", "source_unavailable"} <= kinds


async def test_review_tasks_are_scoped_and_kept_current(db, world):
    task = db.query(WorkflowTask).filter(WorkflowTask.source_id == "workspace:ws-a").one()
    assert db.get(RecordAccess, ("task", task.id)).workspace_id == "ws-a"
    assert "suggested" in task.title


async def test_export_is_scoped_and_formula_safe(db, world):
    from openpyxl import load_workbook

    from services.compliance_graph.export import build

    ada = scope(db, world["ada"])
    review.create_control(db, ada, {"name": "=HYPERLINK(\"http://evil\",\"x\")"})
    review.create_control(db, scope(db, world["bola"]), {"name": "B SECRET CONTROL"})
    data = build(db, ada, generated_by="Ada")
    book = load_workbook(BytesIO(data))
    text = " ".join(str(c.value) for sheet in book.worksheets for row in sheet.iter_rows() for c in row if c.value)
    assert "B SECRET CONTROL" not in text
    assert "Not established in Iroko's records" in text or "not a compliance finding" in text
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"


async def test_iroko_written_text_never_states_a_verdict(db, world):
    ada = scope(db, world["ada"])
    payloads = [queries.overview(db, ada), queries.requirements(db, ada, {}, 1, 100), queries.review_queue(db, ada)]
    text = str(payloads)
    for row in queries.requirement_rows(db, ada):
        assert not wording.has_verdict_wording(row["coverage"]["label"])
        assert not wording.has_verdict_wording(row["applicability"]["label"])
    assert not wording.has_verdict_wording(" ".join(str(link.rationale) for link in db.query(Link)))
    assert "non-compliant" not in text.lower()


async def test_workspaces_never_see_each_others_rows(db, world):
    a, b = scope(db, world["ada"]), scope(db, world["bola"])
    a_control = db.query(Control).filter(Control.workspace_id == "ws-a").first()
    with pytest.raises(GraphNotFound):
        queries.control_detail(db, b, a_control.id)
    a_link = db.query(Link).filter(Link.workspace_id == "ws-a").first()
    with pytest.raises(GraphNotFound):
        queries.link_detail(db, b, a_link.id)
    assert all(c["id"] != a_control.id for c in queries.controls_list(db, b))
    assert "AML training register" not in str(queries.search(db, b, "training"))
    assert not queries.impacts(db, b)
    queue = str(queries.review_queue(db, b))
    assert a_link.id not in queue and "ws-a" not in queue
    with pytest.raises(GraphNotFound):
        queries.neighbourhood(db, b, "control", a_control.id)
    hood = queries.neighbourhood(db, a, "control", a_control.id, depth=2)
    assert any(n["type"] == "requirement" for n in hood["nodes"])
    # Workspace-level counts for B are computed only from what B can see: its own policy's controls.
    overview = queries.overview(db, b)
    b_controls = db.query(Control).filter(Control.workspace_id == "ws-b", Control.review_status == "proposed").count()
    assert overview["totals"]["awaiting_review"] == b_controls == 2
    assert db.query(ObligationStatus).filter(ObligationStatus.workspace_id == "ws-b").count() == 0
    assert core.workspace_links(db, for_workspace("ws-b")) == []
