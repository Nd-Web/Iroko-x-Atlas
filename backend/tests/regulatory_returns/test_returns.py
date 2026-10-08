"""Regulatory returns: deadlines, prudential computations and end-to-end generation."""

from __future__ import annotations

import io
from datetime import date, timedelta

import pytest
from docx import Document
from openpyxl import load_workbook

from services.regulatory_returns import builders, templates
from services.regulatory_returns.calendar import due_date, parse_period, previous_working_day, upcoming
from services.regulatory_returns.catalog import RETURNS, RETURNS_BY_ID
from services.regulatory_returns.prudential import Loan, classify, compute

PROFILE = {
    "institution_name": "Adire Microfinance Bank Limited",
    "licence_category": "state",
    "cbn_licence_no": "MFB/LAG/0412",
    "rc_number": "RC 1234567",
    "head_office_address": "12 Broad Street, Lagos Island, Lagos",
    "md_ceo_name": "Folake Adeyemi",
    "cco_name": "Chinedu Okafor",
    "cfo_name": "Ibrahim Musa",
    "contact_email": "compliance@adiremfb.ng",
    "contact_phone": "+234 801 000 0000",
}

# A balanced statement of financial position for a State MFB (₦).
SFP = {
    "A01": 45_000_000, "A02": 120_000_000, "A03": 210_000_000, "A04": 150_000_000, "A05": 300_000_000,
    "A06": 50_000_000, "A07": 1_600_000_000, "A08": 40_000_000, "A09": 0, "A10": 275_000_000,
    "A11": 180_000_000, "A12": 20_000_000, "A13": 0,
    "L01": 420_000_000, "L02": 610_000_000, "L03": 380_000_000, "L04": 40_000_000,
    "L05": 60_000_000, "L06": 80_000_000, "L07": 12_000_000, "L08": 0, "L09": 48_000_000,
    "E01": 1_000_000_000, "E02": 50_000_000, "E03": 90_000_000, "E04": 0, "E05": 120_000_000, "E06": 0,
}
PL = {"I01": 310_000_000, "I02": 42_000_000, "X01": 95_000_000, "X02": 9_000_000, "I03": 38_000_000, "I04": 6_000_000,
      "X03": 22_000_000, "X04": 110_000_000, "X05": 18_000_000, "X06": 64_000_000, "X07": 23_000_000}


def _loans(total: float = 1_600_000_000) -> list[list]:
    rows = [
        ["B001", "Okon Traders Cooperative", "Cooperative", "N", "", "F001", 60_000_000, 0, "Trade", "N"],
        ["B002", "Amina Bello", "Individual", "N", "", "F002", 15_000_000, 12, "Agriculture", "N"],  # > 1% of SHF -> breach
        ["B003", "Tunde Bakare", "Individual", "Y", "Director", "F003", 8_000_000, 70, "Services", "N"],
        ["B004", "Ngozi Eze", "Individual", "N", "", "F004", 2_000_000, 200, "Trade", "N"],
    ]
    remaining = total - sum(r[6] for r in rows)
    per = remaining / 400
    for i in range(400):
        rows.append([f"C{i:04d}", f"Customer {i}", "Individual", "N", "", f"G{i:04d}", per, 0 if i % 10 else 45, "Trade", "N"])
    return rows


def _prudential_upload(sfp=SFP, pl=PL, loans=None) -> bytes:
    wb = load_workbook(io.BytesIO(templates.prudential_template()))
    for name, values in (("SFP", sfp), ("PL", pl)):
        ws = wb[name]
        for row in ws.iter_rows(min_row=2):
            if row[0].value in values:
                row[2].value = values[row[0].value]
    ws = wb["LoanBook"]
    for r in loans if loans is not None else _loans():
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    for s in doc.sections:
        parts += [p.text for p in s.header.paragraphs]
    return "\n".join(parts)


# ─── Calendar ─────────────────────────────────────────────────────────────────


def test_fina_deadline_rolls_back_over_weekend_and_easter():
    spec = RETURNS_BY_ID["cbn-monthly-prudential"]
    assert due_date(spec, parse_period("month", "2026-09")) == date(2026, 10, 5)  # Monday
    # 5 April 2026 is Easter Sunday; Good Friday 3 April -> Thursday 2 April
    assert due_date(spec, parse_period("month", "2026-03")) == date(2026, 4, 2)
    # 5 July 2026 is a Sunday -> Friday 3 July
    assert due_date(spec, parse_period("month", "2026-06")) == date(2026, 7, 3)


def test_semiannual_and_annual_deadlines():
    assert due_date(RETURNS_BY_ID["cbn-cg-compliance"], parse_period("half_year", "2026-H1")) == date(2026, 7, 14)
    assert due_date(RETURNS_BY_ID["cbn-whistleblowing"], parse_period("half_year", "2026-H2")) == date(2027, 1, 7)
    assert due_date(RETURNS_BY_ID["ndic-deposit-certification"], parse_period("year", "2025")) == date(2026, 1, 31)
    assert due_date(RETURNS_BY_ID["cbn-afs-submission"], parse_period("year", "2025")) == date(2026, 4, 30)
    assert due_date(RETURNS_BY_ID["cbn-afs-submission"], parse_period("year", "2026", fy_end_month=6)) == date(2026, 10, 31)


def test_upcoming_includes_closed_period_with_open_deadline():
    items = upcoming(date(2026, 10, 3))
    first = items[0]
    assert first["return_id"] in ("cbn-monthly-prudential", "cbn-fraud-forgeries")
    assert first["period"] == "2026-09" and first["due"] == "2026-10-05"
    assert all(i["days_left"] >= 0 for i in items)
    assert previous_working_day(date(2026, 10, 1)) == date(2026, 9, 30)  # Independence Day


def test_invalid_period_message():
    with pytest.raises(ValueError, match="2026-09"):
        parse_period("month", "Sept")


# ─── Prudential ───────────────────────────────────────────────────────────────


@pytest.mark.parametrize("dpd,expected", [(0, "Performing"), (30, "Performing"), (31, "Pass & watch"), (59, "Pass & watch"),
                                          (60, "Substandard"), (90, "Substandard"), (91, "Doubtful"), (180, "Doubtful"), (181, "Lost")])
def test_classification_bands(dpd, expected):
    assert classify(dpd)[0] == expected


def test_compute_ratios_and_exposures():
    parsed = templates.parse_prudential(_prudential_upload())
    assert not parsed.problems
    res = compute(parsed.sfp, parsed.pl, parsed.loans, "state")
    t = res.totals
    assert t["total_assets"] == pytest.approx(t["total_liabilities"] + t["equity"])
    assert t["shareholders_funds"] == 1_260_000_000
    liq = next(r for r in res.ratios if r["code"] == "LIQ")
    assert liq["value"] == pytest.approx(825_000_000 / 1_450_000_000)
    assert liq["status"] == "Compliant"
    car = next(r for r in res.ratios if r["code"] == "CAR")
    rwa = 0.2 * (210_000_000 + 150_000_000) + 50_000_000 + 275_000_000 + 1_560_000_000 + 180_000_000
    assert car["value"] == pytest.approx((1_260_000_000 - 20_000_000) / rwa)
    codes = [b.code for b in res.breaches]
    assert codes.count("SOL") == 1  # Amina Bello: 15m > 1% of 1.26bn
    assert "INS" not in codes  # 8m insider < 5% of SHF
    assert any("insider-related exposure" in w for w in res.warnings)  # Tunde Bakare is substandard
    lost = next(c for c in res.classification if c["class"] == "Lost")
    assert lost["count"] == 1 and lost["provision"] == pytest.approx(2_000_000)


def test_unbalanced_sfp_and_reconciliation_warn():
    sfp = dict(SFP, A10=SFP["A10"] + 5_000_000)
    res = compute(sfp, PL, [Loan("X", "X", "Individual", False, "", "F", 1_000_000, 0)], "state")
    assert any("does not balance" in w for w in res.warnings)
    assert any("does not reconcile" in w for w in res.warnings)


def test_capital_breach_detected():
    res = compute(SFP, PL, [], "national")  # ₦1.26bn against the ₦5bn national minimum
    assert any(b.code == "MIN" for b in res.breaches)


def test_template_parse_reports_bad_rows():
    loans = [["B1", "", "Alien", "N", "", "F1", "abc", 5, "", "N"]]
    parsed = templates.parse_prudential(_prudential_upload(loans=loans))
    joined = " ".join(parsed.problems)
    assert "borrower type" in joined and "not a number" in joined and "borrower name is missing" in joined


def test_template_rejects_non_xlsx():
    with pytest.raises(templates.TemplateError):
        templates.parse_prudential(b"not a workbook")


# ─── End to end ───────────────────────────────────────────────────────────────


def _payload(period=None, data=None, remediation=None):
    return {"profile": PROFILE, "period": period, "data": data or {}, "remediation": remediation or {}, "letter_date": "2026-10-03"}


def test_prudential_requires_remediation_then_generates():
    upload = _prudential_upload()
    p = builders.prepare("cbn-monthly-prudential", _payload("2026-09"), upload)
    prev = builders.preview(p)
    assert prev["errors"] == []
    assert not prev["ready"]
    assert [r["code"] for r in prev["remediation_required"]] == ["SOL"]
    with pytest.raises(builders.ReturnInputError):
        builders.render(p)

    p = builders.prepare("cbn-monthly-prudential", _payload("2026-09", remediation={"SOL": "The excess of ₦(...) will be reduced by a partial repayment agreed with the borrower before 31 October 2026."}), upload)
    files = builders.render(p)
    names = [f[0] for f in files]
    assert names == ["AMFB_CBN_MPR_2026-09.docx", "AMFB_CBN_MPR_2026-09.xlsx"]
    text = _docx_text(files[0][1])
    assert "RENDITION OF MONTHLY RETURNS FOR THE MONTH ENDED 30TH SEPTEMBER, 2026" in text
    assert "ADIRE MICROFINANCE BANK LIMITED" in text
    assert "Other Financial Institutions Supervision Department" in text
    assert "Schedule 2 — Statement of financial position" in text
    assert "Single-obligor limit — Amina Bello" in text
    assert "Section 5.2.15" in text
    wb = load_workbook(io.BytesIO(files[1][1]))
    assert wb.sheetnames[:2] == ["MMFBR 300 SFP", "MMFBR 1000 PL (YTD)"]


def test_prudential_without_upload_errors():
    p = builders.prepare("cbn-monthly-prudential", _payload("2026-09"))
    assert any("Upload" in e for e in p.errors)


def test_fraud_nil_and_cases():
    p = builders.prepare("cbn-fraud-forgeries", _payload("2026-09", {"nil_return": True}))
    files = builders.render(p)
    assert len(files) == 1 and "NIL return" in _docx_text(files[0][1])

    case = {"case_ref": "FF/2026/014", "date_occurred": "2026-09-10", "date_discovered": "2026-09-12", "branch": "Ikeja",
            "category": "Cash suppression / theft", "description": "Teller suppressed customer deposits over three days.",
            "amount_involved": 1_250_000, "actual_loss": 400_000, "amount_recovered": 0, "staff_involved": "J. Doe, Teller",
            "reported_to_police": True, "status": "Under investigation", "action_taken": "Staff suspended; matter reported to Ikeja police."}
    p = builders.prepare("cbn-fraud-forgeries", _payload("2026-09", {"nil_return": False, "incidents": [case]}))
    files = builders.render(p)
    text = _docx_text(files[0][1])
    assert "FF/2026/014" in text and "1,250,000.00" in text
    assert len(files) == 2

    bad = builders.prepare("cbn-fraud-forgeries", _payload("2026-09", {"nil_return": True, "incidents": [case]}))
    assert bad.errors


def _cg_data(**over):
    board = [
        {"name": "Adebayo Ogun", "role": "Chairman", "first_appointed": "2019-05-01", "cbn_approved": True, "board_meetings_held": 3, "board_meetings_attended": 3, "committees": ""},
        {"name": "Folake Adeyemi", "role": "MD/CEO", "first_appointed": "2020-01-15", "cbn_approved": True, "board_meetings_held": 3, "board_meetings_attended": 3, "committees": "Credit"},
        {"name": "Kemi Lawal", "role": "Independent Non-Executive Director", "first_appointed": "2022-03-01", "cbn_approved": True, "board_meetings_held": 3, "board_meetings_attended": 3, "committees": "Audit, Governance"},
        {"name": "Yusuf Abdullahi", "role": "Non-Executive Director", "first_appointed": "2019-05-01", "cbn_approved": True, "board_meetings_held": 3, "board_meetings_attended": 1, "committees": "Audit, Risk"},
        {"name": "Grace Nwosu", "role": "Non-Executive Director", "first_appointed": "2021-07-01", "cbn_approved": True, "board_meetings_held": 3, "board_meetings_attended": 3, "committees": "Audit, Credit"},
    ]
    committees = [
        {"name": "Board Risk & Audit Committee (combined)", "chair": "Kemi Lawal", "chair_role": "Independent Non-Executive Director", "members": 3, "meetings_held": 2, "charter_cbn_approved": True},
        {"name": "Board Governance & Nominations Committee", "chair": "Kemi Lawal", "chair_role": "Independent Non-Executive Director", "members": 3, "meetings_held": 2, "charter_cbn_approved": True},
        {"name": "Board Credit Committee", "chair": "Grace Nwosu", "chair_role": "Non-Executive Director", "members": 3, "meetings_held": 1, "charter_cbn_approved": True},
    ]
    data = {"board": board, "committees": committees, "board_meetings_in_period": 3, "external_auditor": "Okoro & Co (Chartered Accountants)",
            "external_auditor_cbn_approved": True, "risk_officer_in_place": True, "internal_auditor_in_place": True,
            "board_appraisal_date": "2026-03-20", "family_members_on_board": False}
    data.update(over)
    return data


def test_cg_matrix_flags_breaches():
    p = builders.prepare("cbn-cg-compliance", _payload("2026-H1", _cg_data()))
    assert p.errors == []
    codes = {b["code"] for b in p.breaches}
    assert codes == {"CG-2.6.1"}  # credit committee met once in the half-year
    statuses = {c["code"]: c["status"] for c in p.ctx["checks"]}
    assert statuses["CG-2.6.3"] == "For information"  # Yusuf attended 1/3
    p = builders.prepare("cbn-cg-compliance", _payload("2026-H1", _cg_data(), {"CG-2.6.1": "The Credit Committee has scheduled meetings for 15 July and 15 September 2026."}))
    files = builders.render(p)
    text = _docx_text(files[0][1])
    assert "Section 8.2" in text and "Compliance matrix" in text and "Credit Committee has scheduled" in text


def test_whistleblowing_nil():
    data = {"policy_approval_date": "2024-02-10", "policy_on_website": True, "channels": "whistle@adiremfb.ng; hotline 0700-ADIRE",
            "cases_received": 0, "cases_investigated": 0, "cases_substantiated": 0, "cases_closed": 0, "cases_pending": 0}
    p = builders.prepare("cbn-whistleblowing", _payload("2026-H1", data))
    files = builders.render(p)
    assert "nil return" in _docx_text(files[0][1])


def test_str_report():
    data = {
        "subject_name": "Bright Star Ventures", "subject_type": "Body corporate", "account_number": "1002003004",
        "detection_datetime": "2026-10-02 14:30",
        "transactions": [{"date": "2026-10-02", "type": "Cash deposit", "amount": 4_900_000, "channel": "Ikeja"},
                         {"date": "2026-10-02", "type": "Outward transfer", "amount": 4_850_000, "counterparty": "Unknown beneficiary"}],
        "grounds": ["(d) Inconsistent with known pattern of the account"],
        "reasons": "Cash deposits structured just below the ₦5,000,000 threshold, immediately transferred out to a newly added beneficiary.",
        "action_taken": "Account placed on enhanced monitoring.", "reporting_officer": "Chinedu Okafor, CCO",
    }
    p = builders.prepare("nfiu-str", {"profile": PROFILE, "data": data})
    assert p.errors == []
    text = _docx_text(builders.render(p)[0][1])
    assert "SUSPICIOUS TRANSACTION REPORT" in text and "19(1)(a)" in text and "9,750,000.00" in text


def test_ctr_selects_only_amounts_above_threshold():
    wb = load_workbook(io.BytesIO(templates.ctr_template()))
    ws = wb["Transactions"]
    today = date.today()
    ws.append([today.isoformat(), "Ada Obi", "Individual", "001", "222", "Cash deposit", 5_000_000, "", "Yaba", ""])       # equal — not "in excess"
    ws.append([today.isoformat(), "Ada Obi", "Individual", "001", "222", "Cash deposit", 5_000_001, "", "Yaba", ""])       # reportable
    ws.append([today.isoformat(), "Zenon Ltd", "Body corporate", "002", "RC9", "Cash withdrawal", 9_000_000, "", "Yaba", ""])  # below corporate
    ws.append([(today - timedelta(days=10)).isoformat(), "Zenon Ltd", "Body corporate", "002", "RC9", "Cash deposit", 12_000_000, "", "Yaba", ""])  # late
    buf = io.BytesIO()
    wb.save(buf)
    data = {"period_from": (today - timedelta(days=14)).isoformat(), "period_to": today.isoformat(), "reporting_officer": "Chinedu Okafor"}
    p = builders.prepare("nfiu-ctr", {"profile": PROFILE, "data": data}, buf.getvalue())
    assert p.errors == []
    assert len(p.ctx["rows"]) == 2
    assert sum(1 for r in p.ctx["rows"] if r["late"]) == 1
    assert any("7-day" in w for w in p.warnings)
    assert len(builders.render(p)) == 2


def test_ndic_afs_ndpc_generate():
    ndic = {"demand_deposits": 420e6, "savings_deposits": 610e6, "time_deposits": 380e6, "other_deposits": 40e6,
            "insider_deposits": 15e6, "collateral_deposits": 35e6, "auditor_name": "Okoro & Co", "auditor_certificate_date": "2026-01-20"}
    p = builders.prepare("ndic-deposit-certification", _payload("2025", ndic))
    text = _docx_text(builders.render(p)[0][1])
    assert "1,450,000,000.00" in text and "1,400,000,000.00" in text and "Section 17" in text

    afs = {"auditor_name": "Okoro & Co", "audit_report_date": "2026-03-28", "board_approval_date": "2026-03-25",
           "total_assets": 2.6e9, "profit_before_tax": 85e6, "shareholders_funds": 1.26e9, "management_letter_enclosed": True}
    p = builders.prepare("cbn-afs-submission", _payload("2025", afs))
    assert "not publish" in _docx_text(builders.render(p)[0][1])

    ndpc = {"dpco_name": "SecureData DPCO Ltd", "dpo_name": "Halima Sani", "data_subjects": 48000, "data_categories": "Identity, contact, financial",
            "lawful_bases": "Contract; legal obligation", "dpias_conducted": 2, "breaches": 1, "breaches_notified_72h": 1,
            "dsr_received": 14, "dsr_resolved": 14, "security_measures": "Encryption at rest, MFA, quarterly access reviews"}
    p = builders.prepare("ndpc-compliance-audit", _payload("2025", ndpc))
    assert "31st December, 2025" in _docx_text(builders.render(p)[0][1])


def test_required_fields_and_profile_validation():
    p = builders.prepare("cbn-whistleblowing", {"profile": {}, "period": "2026-H1", "data": {}})
    assert any("Institution name is required" in e for e in p.errors)
    assert any("Reporting channels is required" in e for e in p.errors)


def test_every_generator_is_wired():
    for spec in RETURNS:
        if spec.generator:
            assert spec.id in builders._COMPUTE and spec.id in builders._RENDER and spec.id in builders.REF_CODES


def test_unbalanced_balance_sheet_blocks_generation():
    sfp = dict(SFP, A10=SFP["A10"] + 9_000_000)
    p = builders.prepare("cbn-monthly-prudential", _payload("2026-09"), dataset={"sfp": sfp, "pl": PL, "loans": []})
    assert any("does not balance" in e for e in p.errors)
    with pytest.raises(builders.ReturnInputError):
        builders.render(p)
