"""
services/regulatory_returns/importers.py

Read the files a bank already produces — trial balance, loan book, cash
transactions — in whatever layout its core banking system exports, and turn
them into return data.

  * Columns are found by header synonyms; when a required column is still
    missing, the model is asked to pick it from the real headers, and the
    officer can always override the choice.
  * Trial-balance accounts map to MMFBR lines in this order: the bank's own
    mapping from earlier months (exact GL code), accounting keyword rules, then
    the model. Every account shows how it was mapped, and an unmapped account
    is never guessed silently.
  * The officer's corrections are saved back to memory, so next month the
    whole trial balance maps without help.

The model is only ever asked to classify labels it is shown; every amount
comes from the file.
"""

from __future__ import annotations

import csv
import io
import json
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Awaitable, Callable

from openpyxl import load_workbook

from .prudential import BORROWER_TYPES, PL_LINES, SFP_LINES, Loan
from .templates import MAX_UPLOAD_BYTES, TemplateError, Txn

logger = logging.getLogger(__name__)

MAX_ROWS = 50_000
LINE_LABELS = {code: label for code, label, _ in SFP_LINES + PL_LINES}
LINE_CODES = tuple(LINE_LABELS) + ("NONE",)

Complete = Callable[..., Awaitable[str]]

# ─── Reading any table ────────────────────────────────────────────────────────


def _norm(text: Any) -> str:
    t = str(text or "").strip().lower()
    t = re.sub(r"\(.*?\)|₦|ngn|naira", " ", t)
    return " ".join(re.sub(r"[^a-z0-9/ ]", " ", t).split())


@dataclass
class Table:
    headers: list[str]
    rows: list[list[Any]]
    sheet: str = ""


def _cells(row) -> list[Any]:
    return [c.isoformat() if isinstance(c, (datetime, date)) else c for c in row]


def read_table(data: bytes, filename: str, synonyms: dict[str, tuple[str, ...]]) -> Table:
    """The first sheet/table whose header row matches the most known columns."""
    if not data:
        raise TemplateError(["The uploaded file is empty."])
    if len(data) > MAX_UPLOAD_BYTES:
        raise TemplateError(["The uploaded file exceeds the 10 MB limit."])
    name = (filename or "").lower()
    candidates: list[tuple[str, list[list[Any]]]] = []
    if name.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        candidates.append(("csv", [r for r in csv.reader(io.StringIO(text), dialect)]))
    else:
        try:
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except Exception:
            raise TemplateError(["Iroko could not read this file. Upload the Excel (.xlsx) or CSV export from your core banking system."]) from None
        for ws in wb.worksheets:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS:
                    break
                rows.append(_cells(r))
            candidates.append((ws.title, rows))

    wanted = {_norm(s) for group in synonyms.values() for s in group}
    best: tuple[int, str, int, list[list[Any]]] | None = None
    for sheet, rows in candidates:
        for idx, row in enumerate(rows[:30]):
            score = sum(1 for c in row if _norm(c) in wanted)
            if score >= 2 and (best is None or score > best[0]):
                best = (score, sheet, idx, rows)
    if best is None:
        raise TemplateError(["Iroko could not find a header row in this file. Make sure the export includes column headings."])
    _, sheet, idx, rows = best
    headers = [str(c).strip() if c is not None else "" for c in rows[idx]]
    body = [r for r in rows[idx + 1:] if any(v not in (None, "") for v in r)]
    return Table(headers, body, sheet)


def detect_columns(headers: list[str], synonyms: dict[str, tuple[str, ...]]) -> dict[str, int]:
    found: dict[str, int] = {}
    normed = [_norm(h) for h in headers]
    # exact matches first, then "contains", so "balance" doesn't steal "loan balance"
    for exact in (True, False):
        for field_key, syns in synonyms.items():
            if field_key in found:
                continue
            for s in syns:
                ns = _norm(s)
                for i, h in enumerate(normed):
                    if i in found.values() or not h:
                        continue
                    if (h == ns) if exact else (ns in h.split() or (len(ns) > 3 and ns in h)):
                        found[field_key] = i
                        break
                if field_key in found:
                    break
    return found


_COLUMN_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["columns"],
    "properties": {"columns": {"type": "array", "items": {
        "type": "object", "additionalProperties": False, "required": ["field", "header"],
        "properties": {"field": {"type": "string"}, "header": {"type": "string"}},
    }}},
}


async def ai_columns(table: Table, missing: list[str], descriptions: dict[str, str], complete: Complete | None) -> dict[str, int]:
    if not complete or not missing:
        return {}
    sample = [table.headers] + table.rows[:5]
    prompt = (
        "A Nigerian microfinance bank exported this table from its core banking system.\n"
        f"First rows (JSON):\n{json.dumps(sample, default=str)[:6000]}\n\n"
        "For each field below, choose the ONE header from the first row that holds it, copied exactly. "
        "If no column holds it, use an empty string.\n"
        + "\n".join(f"- {k}: {descriptions[k]}" for k in missing)
    )
    try:
        raw = await complete(prompt, json_schema=_COLUMN_SCHEMA, max_tokens=400, service_id="nano")
        picks = json.loads(raw or "{}").get("columns", [])
    except Exception as exc:  # the model is optional; the officer can still choose
        logger.warning("[importers] column detection by model failed: %s", exc)
        return {}
    out = {}
    for p in picks:
        if p.get("field") in missing and p.get("header") in table.headers:
            out[p["field"]] = table.headers.index(p["header"])
    return out


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("₦", "").replace("NGN", "").strip()
    neg = text.startswith("(") and text.endswith(")") or text.endswith("-")
    text = text.strip("()").rstrip("-").strip()
    if text.upper().endswith(("DR", "CR")):
        sign = -1 if text.upper().endswith("CR") else 1
        text = text[:-2].strip()
    else:
        sign = 1
    try:
        n = float(text)
    except ValueError:
        return None
    return -n * sign if neg else n * sign


# ─── Trial balance ────────────────────────────────────────────────────────────

TB_SYNONYMS: dict[str, tuple[str, ...]] = {
    "code": ("gl code", "account code", "gl account", "gl no", "gl number", "ledger code", "account no", "account number", "acct code", "code", "gl"),
    "name": ("account name", "gl name", "ledger name", "account description", "description", "account title", "particulars", "name"),
    "debit": ("debit", "debit balance", "dr", "debit amount", "closing debit"),
    "credit": ("credit", "credit balance", "cr", "credit amount", "closing credit"),
    "balance": ("closing balance", "net balance", "balance", "ending balance", "closing bal", "amount"),
    "drcr": ("dr/cr", "dr cr", "balance type", "sign"),
}
TB_DESCRIPTIONS = {
    "code": "the GL / account code", "name": "the account name or description",
    "debit": "the debit balance", "credit": "the credit balance", "balance": "a single signed or unsigned closing balance",
}

# Keyword rules: (pattern, line, requires) — requires "dr"/"cr" restricts by balance side.
_RULES: tuple[tuple[str, str, str], ...] = (
    (r"accumulated (depreciation|amorti)", "A11", ""),
    (r"(impairment|provision).*(loan|advance|credit)|loan loss provision|allowance for", "A08", "cr"),
    (r"regulatory risk reserve", "E04", ""),
    (r"statutory reserve", "E03", ""),
    (r"share premium", "E02", ""),
    (r"share capital|paid.?up capital|ordinary shares", "E01", ""),
    (r"retained (earnings|profit)|accumulated (loss|profit)|revenue reserve", "E05", ""),
    (r"deposit for shares|other reserve|revaluation reserve", "E06", ""),
    (r"cash (in|on) (hand|vault|till)|vault cash|teller cash|petty cash|^cash$", "A01", ""),
    (r"central bank|\bcbn\b", "A02", "dr"),
    (r"treasury bill|fgn bond|government securit|\bt-?bill", "A05", "dr"),
    (r"placement|fixed deposit with|call deposit with|money market", "A04", "dr"),
    (r"due from|balances? with (other )?(banks?|ofis?)|bank balances?|nostro|current accounts? with", "A03", "dr"),
    # Accounts named after a Nigerian bank (settlement / current accounts held with it)
    (r"\b(zenith|access|gtbank|gtb|guaranty trust|first ?bank|uba|united bank for africa|wema|fidelity|sterling|stanbic|"
     r"union bank|polaris|ecobank|fcmb|first city|keystone|unity bank|providus|titan trust|globus|optimus|parallex|"
     r"premium trust|signature bank|lotus|jaiz|taj ?bank|suntrust|heritage bank)\b", "A03", "dr"),
    (r"intangible|software|goodwill", "A12", "dr"),
    (r"deferred tax asset", "A13", "dr"),
    (r"deferred tax liabilit", "L08", "cr"),
    (r"investment in (subsidiar|associate)", "A09", "dr"),
    (r"land|building|motor vehicle|furniture|fitting|office equipment|computer equipment|plant|machinery|leasehold|property", "A11", ""),
    (r"interest (income|earned).*(loan|advance|overdraft)|loan interest income|interest on loans", "I01", "cr"),
    (r"interest (income|earned).*(placement|treasury|investment|bank)|investment income", "I02", "cr"),
    (r"interest (expense|paid).*(deposit|savings|fixed|term)|interest on deposit", "X01", "dr"),
    (r"(interest|finance) (expense|cost|charge|paid|on)\b.*(borrow|loans? from|facilit|on.?lending|takings)|finance cost", "X02", "dr"),
    (r"fee|commission|charges income|account maintenance", "I03", "cr"),
    (r"impairment (charge|expense)|loan loss expense|provision (charge|expense)|bad debt", "X03", "dr"),
    (r"salar|wage|staff|pension|allowance|bonus|payroll|gratuity", "X04", "dr"),
    (r"depreciation|amortisation|amortization", "X05", "dr"),
    (r"income tax expense|tax expense|education tax|company income tax charge", "X07", "dr"),
    (r"(current )?income tax payable|tax payable|cit payable", "L07", "cr"),
    (r"mandatory saving|cash collateral|compulsory saving|other deposit", "L04", "cr"),
    (r"current account|demand deposit|current deposit", "L01", "cr"),
    (r"savings", "L02", "cr"),
    (r"fixed deposit|term deposit|time deposit", "L03", "cr"),
    (r"borrowing from (bank|ofi)|takings|interbank", "L05", "cr"),
    (r"on.?lending|intervention fund|borrowing|loan from", "L06", "cr"),
    (r"loan|advance|overdraft|lease receivable|microcredit|credit facilit", "A07", "dr"),
    (r"prepayment|receivable|sundry debtor|other asset|account receivable|stock|inventory", "A10", "dr"),
    (r"accru|payable|sundry creditor|other liabilit|unearned|deferred income|provision for", "L09", "cr"),
    (r"other (operating )?income|recover|gain on|rental income|sundry income", "I04", "cr"),
    (r"expense|rent|utilit|maintenance|repairs|insurance|audit fee|legal|advert|transport|travel|telephone|internet|printing|stationery|fuel|diesel|security|cleaning|training|donation|bank charges|directors.? fee|subscription", "X06", "dr"),
)


def rule_line(name: str, net: float) -> str | None:
    n = name.lower()
    side = "dr" if net > 0 else "cr" if net < 0 else ""
    for pattern, line, requires in _RULES:
        if re.search(pattern, n) and (not requires or not side or requires == side):
            return line
    return None


_MAP_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["mappings"],
    "properties": {"mappings": {"type": "array", "items": {
        "type": "object", "additionalProperties": False, "required": ["code", "line"],
        "properties": {"code": {"type": "string"}, "line": {"type": "string", "enum": list(LINE_CODES)}},
    }}},
}


async def ai_lines(accounts: list[dict], complete: Complete | None) -> dict[str, str]:
    if not complete or not accounts:
        return {}
    lines_text = "\n".join(f"{c}: {l}" for c, l in LINE_LABELS.items())
    out: dict[str, str] = {}
    for start in range(0, len(accounts), 120):
        batch = accounts[start:start + 120]
        listing = "\n".join(
            f"{a['code']} | {a['name']} | {'debit' if a['net'] > 0 else 'credit' if a['net'] < 0 else 'nil'} balance"
            for a in batch
        )
        prompt = (
            "Map each general-ledger account of a Nigerian microfinance bank to one CBN MMFBR line "
            "(statement of financial position A/L/E lines, profit or loss I/X lines). Use NONE only for "
            "memorandum, suspense or contra-memo accounts that belong on no line. A02 is ONLY money held at the "
            "Central Bank of Nigeria; settlement or current accounts with commercial banks are A03. Debit-balance "
            "accounts are normally assets (A) or expenses (X); credit-balance accounts are liabilities (L), equity (E) "
            "or income (I). Income and expense accounts are never NONE.\n\n"
            f"Lines:\n{lines_text}\nNONE: not reported\n\nAccounts (code | name | balance side):\n{listing}"
        )
        try:
            raw = await complete(prompt, json_schema=_MAP_SCHEMA, max_tokens=4000, service_id="nano")
            for m in json.loads(raw or "{}").get("mappings", []):
                if m.get("line") in LINE_CODES:
                    out[str(m.get("code"))] = m["line"]
        except Exception as exc:
            logger.warning("[importers] GL mapping by model failed: %s", exc)
            break
    return out


def _line_value(line: str, net: float, signed: bool, name: str) -> float:
    credit_lines = line.startswith(("L", "E", "I")) or line == "A08"
    if signed:
        return -net if credit_lines else net
    v = abs(net)
    if line == "A11" and re.search(r"accumulated", name.lower()):
        return -v
    return v


def compute_tb(accounts: list[dict], signed: bool) -> dict:
    """Totals per MMFBR line from mapped accounts, plus current-year profit."""
    sfp = {c: 0.0 for c, _, _ in SFP_LINES}
    pl = {c: 0.0 for c, _, _ in PL_LINES}
    for a in accounts:
        line = a.get("line")
        if not line or line == "NONE":
            continue
        v = _line_value(line, a["net"], signed, a["name"])
        (sfp if line in sfp else pl)[line] += v
    income = sum(pl[c] for c in ("I01", "I02", "I03", "I04"))
    expense = sum(pl[c] for c in ("X01", "X02", "X03", "X04", "X05", "X06", "X07"))
    profit = income - expense
    sfp["E05"] += profit
    notes = [f"Current-year profit of ₦{profit:,.2f} (income less expenses in the trial balance) is included in retained earnings (E05)."]
    warnings = []
    if signed:
        total = sum(a["net"] for a in accounts)
        if abs(total) > 1.0:
            warnings.append(f"The trial balance does not balance: debits exceed credits by ₦{total:,.2f}. Check the export is complete.")
    else:
        warnings.append("The balances in this file have no debit/credit sign, so Iroko applied the natural side of each line. Confirm the statement of financial position balances.")
    unmapped = [a for a in accounts if not a.get("line")]
    return {"sfp": sfp, "pl": pl, "notes": notes, "warnings": warnings, "unmapped": len(unmapped)}


async def import_trial_balance(data: bytes, filename: str, memory: dict[str, str], complete: Complete | None, columns: dict[str, int] | None = None) -> dict:
    table = read_table(data, filename, TB_SYNONYMS)
    cols = dict(columns or detect_columns(table.headers, TB_SYNONYMS))
    if "name" not in cols:
        cols.update(await ai_columns(table, ["name"], TB_DESCRIPTIONS, complete))
    if not (("debit" in cols and "credit" in cols) or "balance" in cols):
        cols.update(await ai_columns(table, [k for k in ("debit", "credit", "balance") if k not in cols], TB_DESCRIPTIONS, complete))
    problems = []
    if "name" not in cols:
        problems.append("account name")
    if not (("debit" in cols and "credit" in cols) or "balance" in cols):
        problems.append("debit and credit (or balance)")
    if problems:
        return {"kind": "trial_balance", "file_name": filename, "headers": table.headers, "columns": cols,
                "needs_columns": problems, "accounts": [], "rows": table.rows[:MAX_ROWS]}
    result = await _build_tb(filename, table.headers, table.rows, cols, memory, complete)
    result["rows"] = table.rows[:MAX_ROWS]
    return result


async def rebuild_tb(dataset: dict, columns: dict[str, int], memory: dict[str, str], complete: Complete | None) -> dict:
    """Re-read stored rows with columns the officer chose."""
    result = await _build_tb(dataset["file_name"], dataset["headers"], dataset["rows"], columns, memory, complete)
    result["rows"] = dataset["rows"]
    return result


async def _build_tb(filename: str, headers: list[str], rows: list[list[Any]], cols: dict[str, int], memory: dict[str, str], complete: Complete | None) -> dict:
    get = lambda r, k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
    accounts: list[dict] = []
    signed = "debit" in cols and "credit" in cols or "drcr" in cols
    any_negative = False
    for r in rows:
        name = str(get(r, "name") or "").strip()
        code = str(get(r, "code") or "").strip()
        if not name and not code:
            continue
        if not code and re.search(r"\btotal\b", name.lower()):
            continue  # subtotal rows
        if "debit" in cols and "credit" in cols:
            net = (_num(get(r, "debit")) or 0.0) - (_num(get(r, "credit")) or 0.0)
        else:
            net = _num(get(r, "balance")) or 0.0
            if "drcr" in cols and str(get(r, "drcr") or "").strip().lower().startswith("c"):
                net = -abs(net)
            elif "drcr" in cols:
                net = abs(net)
            any_negative = any_negative or net < 0
        accounts.append({"code": code or name, "name": name or code, "net": round(net, 2)})
    signed = signed or any_negative
    for a in accounts:
        if a["code"] in memory:
            a.update(line=memory[a["code"]], how="Your mapping from a previous month")
            continue
        line = rule_line(a["name"], a["net"] if signed else 0.0)
        if line:
            a.update(line=line, how="Accounting rule")
    pending = [a for a in accounts if not a.get("line")]
    for code, line in (await ai_lines(pending, complete)).items():
        for a in pending:
            if a["code"] == code:
                a.update(line=line, how="Iroko AI — please confirm")
    for a in accounts:
        a.setdefault("line", None)
        a.setdefault("how", "Not mapped — choose a line")
    result = {"kind": "trial_balance", "file_name": filename, "headers": headers, "columns": cols, "signed": signed,
              "needs_columns": [], "accounts": accounts}
    result.update(compute_tb(accounts, signed))
    return result


def remap_tb(dataset: dict, changes: dict[str, str]) -> dict:
    for a in dataset.get("accounts", []):
        if a["code"] in changes:
            line = changes[a["code"]]
            a["line"] = line if line in LINE_CODES else None
            a["how"] = "Set by you"
    dataset.update(compute_tb(dataset["accounts"], dataset.get("signed", True)))
    return dataset


# ─── Loan book ────────────────────────────────────────────────────────────────

LOAN_SYNONYMS: dict[str, tuple[str, ...]] = {
    "borrower_id": ("customer id", "customer no", "customer number", "cif", "cif no", "borrower id", "client id", "client no", "customer code"),
    "borrower_name": ("customer name", "borrower name", "borrower", "client name", "account name", "name"),
    "facility_id": ("loan id", "loan number", "loan no", "loan account", "facility id", "loan ref", "account number", "loan account number"),
    "outstanding": ("outstanding balance", "principal outstanding", "loan balance", "total outstanding", "amount outstanding", "outstanding", "exposure", "balance"),
    "dpd": ("days past due", "dpd", "days in arrears", "arrears days", "days overdue", "overdue days", "days delinquent", "aging days"),
    "borrower_type": ("borrower type", "customer type", "client type", "customer category", "category", "segment"),
    "insider": ("insider", "related party", "is staff", "staff loan", "insider (y/n)", "insider y/n"),
    "relationship": ("insider relationship", "relationship"),
    "sector": ("sector", "industry", "economic sector", "purpose"),
    "restructured": ("restructured", "rescheduled", "restructured (y/n)"),
}
LOAN_DESCRIPTIONS = {
    "borrower_name": "the borrower or customer name", "outstanding": "the outstanding loan balance in naira",
    "dpd": "the number of days the loan is past due / in arrears", "borrower_id": "a customer identifier (CIF, customer number)",
}
LOAN_REQUIRED = ("borrower_name", "outstanding", "dpd")

_CORP = re.compile(r"\b(ltd|limited|plc|enterprises?|ventures?|nig|nigeria|company|co\.|& sons|industries|services|global|resources|concepts?|integrated|investment)\b", re.I)
_COOP = re.compile(r"\b(co-?operative|coop|cooperative|cms|thrift|society)\b", re.I)
_GROUP = re.compile(r"\b(group|association|union|club|women|farmers|traders)\b", re.I)


def infer_type(name: str, raw: Any) -> tuple[str, bool]:
    """(type, inferred?)"""
    text = str(raw or "").strip().lower()
    for t in BORROWER_TYPES:
        if text.startswith(t.lower()[:4]):
            return t, False
    if text in ("corp", "company", "sme", "business", "corporate"):
        return "Corporate", False
    if _COOP.search(name):
        return "Cooperative", True
    if _GROUP.search(name):
        return "Group", True
    if _CORP.search(name):
        return "Corporate", True
    return "Individual", True


def _yes(v: Any) -> bool:
    return str(v or "").strip().lower() in {"y", "yes", "true", "1", "insider", "staff", "director"}


def _name_key(name: str) -> set[str]:
    return {w for w in re.sub(r"[^a-z ]", " ", name.lower()).split() if len(w) > 2}


async def import_loan_book(data: bytes, filename: str, insiders: list[str], complete: Complete | None, columns: dict[str, int] | None = None) -> dict:
    table = read_table(data, filename, LOAN_SYNONYMS)
    return await build_loan_book(filename, table.headers, table.rows, insiders, complete, columns)


async def build_loan_book(filename: str, headers: list[str], rows: list[list[Any]], insiders: list[str], complete: Complete | None, columns: dict[str, int] | None = None) -> dict:
    table = Table(headers, rows)
    cols = dict(columns or detect_columns(headers, LOAN_SYNONYMS))
    missing = [k for k in LOAN_REQUIRED if k not in cols]
    if missing:
        cols.update(await ai_columns(table, missing, LOAN_DESCRIPTIONS, complete))
    missing = [k for k in LOAN_REQUIRED if k not in cols]
    base = {"kind": "loan_book", "file_name": filename, "headers": headers, "columns": cols, "rows": rows[:MAX_ROWS]}
    if missing:
        names = {"borrower_name": "borrower name", "outstanding": "outstanding balance", "dpd": "days past due"}
        return {**base, "needs_columns": [names[m] for m in missing], "loans": []}
    get = lambda r, k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
    insider_keys = [(n, _name_key(n)) for n in insiders if n]
    loans, inferred_types, matched = [], 0, []
    for i, r in enumerate(rows, start=2):
        name = str(get(r, "borrower_name") or "").strip()
        amount = _num(get(r, "outstanding"))
        if not name or amount is None or re.search(r"\btotal\b", name.lower()):
            continue
        btype, guessed = infer_type(name, get(r, "borrower_type"))
        inferred_types += guessed
        insider = _yes(get(r, "insider")) if "insider" in cols else False
        relationship = str(get(r, "relationship") or "").strip()
        if not insider:
            keys = _name_key(name)
            for full, k in insider_keys:
                if len(k) >= 2 and k <= keys:
                    insider, relationship = True, relationship or f"Director ({full})"
                    matched.append(name)
                    break
        loans.append({
            "borrower_id": str(get(r, "borrower_id") or "").strip(), "borrower_name": name, "borrower_type": btype,
            "insider": insider, "insider_relationship": relationship, "facility_id": str(get(r, "facility_id") or "").strip(),
            "outstanding": round(abs(amount), 2), "days_past_due": max(int(_num(get(r, "dpd")) or 0), 0),
            "sector": str(get(r, "sector") or "").strip(), "restructured": _yes(get(r, "restructured")), "row": i,
        })
    counts = {t: sum(1 for l in loans if l["borrower_type"] == t) for t in BORROWER_TYPES}
    notes = [f"{len(loans)} loans totalling ₦{sum(l['outstanding'] for l in loans):,.2f}."]
    if inferred_types:
        notes.append("Borrower types were inferred from names: " + ", ".join(f"{v} {k.lower()}" for k, v in counts.items() if v) + ". Correct any that are wrong.")
    if matched:
        notes.append(f"{len(matched)} borrower(s) match directors on your board list and are treated as insiders: {', '.join(matched[:5])}{'…' if len(matched) > 5 else ''}.")
    elif "insider" not in cols:
        notes.append("The file has no insider column and no borrower matched your board list. Mark insiders below if any.")
    return {**base, "needs_columns": [], "loans": loans, "notes": notes, "matched_insiders": matched}


def loans_from(dataset: dict) -> list[Loan]:
    return [Loan(**{k: v for k, v in l.items()}) for l in dataset.get("loans", [])]


# ─── Cash transactions (CTR) ─────────────────────────────────────────────────

TXN_SYNONYMS: dict[str, tuple[str, ...]] = {
    "date": ("transaction date", "trans date", "tran date", "value date", "posting date", "txn date", "date"),
    "customer_name": ("customer name", "account name", "customer", "name"),
    "customer_type": ("customer type", "account type", "customer category", "category", "segment"),
    "account_number": ("account number", "account no", "acct no", "nuban", "account"),
    "bvn_or_rc": ("bvn / rc number", "bvn", "rc number", "bvn/rc", "id number", "tin"),
    "transaction_type": ("transaction type", "tran type", "txn type", "type"),
    "amount": ("transaction amount", "amount", "value", "tran amount"),
    "debit": ("debit", "withdrawal", "dr amount"),
    "credit": ("credit", "deposit", "lodgement", "cr amount"),
    "narration": ("narration", "description", "remarks", "details", "particulars"),
    "branch": ("branch", "branch name", "sol"),
    "channel": ("channel", "teller", "teller id", "posted by", "user"),
}
TXN_DESCRIPTIONS = {"date": "the transaction date", "customer_name": "the customer's name", "amount": "the transaction amount in naira"}


def _to_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = str(v or "").strip()
    if not text:
        return None
    if re.match(r"^\d{4}-\d{2}-\d{2}T", text):  # ISO datetime ("04-OCT-2026" also contains a T)
        text = text[:10]
    text = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?.*$", "", text)  # drop a trailing time
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%Y/%m/%d", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


async def import_transactions(data: bytes, filename: str, complete: Complete | None, columns: dict[str, int] | None = None) -> dict:
    table = read_table(data, filename, TXN_SYNONYMS)
    return await build_transactions(filename, table.headers, table.rows, complete, columns)


async def build_transactions(filename: str, headers: list[str], rows: list[list[Any]], complete: Complete | None, columns: dict[str, int] | None = None) -> dict:
    table = Table(headers, rows)
    cols = dict(columns or detect_columns(table.headers, TXN_SYNONYMS))
    need = [k for k in ("date", "customer_name") if k not in cols]
    if "amount" not in cols and not ("debit" in cols or "credit" in cols):
        need.append("amount")
    if need:
        cols.update(await ai_columns(table, need, TXN_DESCRIPTIONS, complete))
    missing = [k for k in ("date", "customer_name") if k not in cols]
    if "amount" not in cols and not ("debit" in cols or "credit" in cols):
        missing.append("amount")
    base = {"kind": "transactions", "file_name": filename, "headers": table.headers, "columns": cols, "rows": table.rows[:MAX_ROWS]}
    if missing:
        names = {"date": "transaction date", "customer_name": "customer name", "amount": "amount"}
        return {**base, "needs_columns": [names[m] for m in missing], "txns": []}
    get = lambda r, k: r[cols[k]] if k in cols and cols[k] < len(r) else None  # noqa: E731
    txns, bad_dates, inferred = [], 0, 0
    for i, r in enumerate(table.rows, start=2):
        name = str(get(r, "customer_name") or "").strip()
        if not name:
            continue
        d = _to_date(get(r, "date"))
        if d is None:
            bad_dates += 1
            continue
        amount = _num(get(r, "amount")) if "amount" in cols else max(_num(get(r, "debit")) or 0.0, _num(get(r, "credit")) or 0.0)
        if amount is None:
            continue
        raw_type = str(get(r, "customer_type") or "").lower()
        if raw_type.startswith(("ind", "per", "sav")):
            ctype = "Individual"
        elif raw_type.startswith(("corp", "bus", "body", "com", "sme")):
            ctype = "Body corporate"
        else:
            ctype = "Body corporate" if _CORP.search(name) or _COOP.search(name) else "Individual"
            inferred += 1
        txns.append({
            "date": d.isoformat(), "customer_name": name, "customer_type": ctype,
            "account_number": str(get(r, "account_number") or "").strip(), "bvn_or_rc": str(get(r, "bvn_or_rc") or "").strip(),
            "transaction_type": str(get(r, "transaction_type") or "").strip(), "amount": round(abs(amount), 2),
            "narration": str(get(r, "narration") or "").strip(), "branch": str(get(r, "branch") or "").strip(),
            "channel": str(get(r, "channel") or "").strip(), "row": i,
        })
    notes = [f"{len(txns)} transactions read."]
    if inferred:
        notes.append(f"Customer type was inferred from the name for {inferred} transaction(s); company-style names were treated as body corporate.")
    if bad_dates:
        notes.append(f"{bad_dates} row(s) had no readable date and were skipped.")
    return {**base, "needs_columns": [], "txns": txns, "notes": notes}


def txns_from(dataset: dict) -> list[Txn]:
    return [Txn(**{**t, "date": date.fromisoformat(t["date"])}) for t in dataset.get("txns", [])]
