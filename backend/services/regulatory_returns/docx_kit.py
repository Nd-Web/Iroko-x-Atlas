"""
services/regulatory_returns/docx_kit.py

Shared layout for regulatory correspondence: the institution's letterhead,
formal letter blocks (reference, date, addressee, salutation, subject),
schedules, certification and signature blocks, and a footer with page
numbers. Returns carry the bank's identity only — no Iroko branding — because
they go out under the bank's name.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Arial"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def naira(value: float | None, dash_zero: bool = False) -> str:
    if value is None:
        return "—"
    if dash_zero and abs(value) < 0.005:
        return "—"
    return f"({abs(value):,.2f})" if value < 0 else f"{value:,.2f}"


def pct(value: float | None, places: int = 2) -> str:
    return "—" if value is None else f"{value * 100:.{places}f}%"


def long_date(d: date) -> str:
    day = d.day
    suffix = "th" if 11 <= day <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix} {d.strftime('%B, %Y')}"


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _page_field(paragraph) -> None:
    for instr in ("PAGE",):
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        text = OxmlElement("w:instrText")
        text.set(qn("xml:space"), "preserve")
        text.text = instr
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(text)
        run._r.append(end)


class ReturnDocument:
    """A letter-format regulatory return under the institution's letterhead."""

    def __init__(self, profile: dict, reference: str, landscape: bool = False):
        self.profile = profile
        self.reference = reference
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        section = self.doc.sections[0]
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.2))
        section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
        self._letterhead()
        self._footer()

    # ── Page furniture ────────────────────────────────────────────────────────

    def _letterhead(self) -> None:
        header = self.doc.sections[0].header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(self.profile.get("institution_name", "").upper())
        run.bold, run.font.size, run.font.color.rgb = True, Pt(14), NAVY
        details = [self.profile.get("head_office_address", "").replace("\n", ", ")]
        ids = []
        if self.profile.get("cbn_licence_no"):
            ids.append(f"CBN Licence No: {self.profile['cbn_licence_no']}")
        rc = (self.profile.get("rc_number") or "").strip()
        if rc:
            ids.append(rc if rc.upper().startswith("RC") else f"RC {rc}")
        contact = " · ".join(x for x in (self.profile.get("contact_phone"), self.profile.get("contact_email")) if x)
        for line in (details[0], " · ".join(ids), contact):
            if line:
                q = header.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.paragraph_format.space_after = Pt(0)
                r = q.add_run(line)
                r.font.size = Pt(8.5)
        rule = header.add_paragraph()
        p_pr = rule._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(k), v)
        border.append(bottom)
        p_pr.append(border)

    def _footer(self) -> None:
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{self.profile.get('institution_name', '')} · Ref: {self.reference} · Confidential · Page ")
        r.font.size = Pt(8)
        _page_field(p)

    # ── Letter blocks ─────────────────────────────────────────────────────────

    def letter_opening(self, letter_date: date, addressee: Sequence[str], subject: str, attention: str | None = None) -> None:
        p = self.doc.add_paragraph()
        p.add_run(f"Our Ref: {self.reference}").bold = True
        p.paragraph_format.space_after = Pt(0)
        self.doc.add_paragraph(long_date(letter_date))
        for i, line in enumerate(addressee):
            q = self.doc.add_paragraph(line)
            q.paragraph_format.space_after = Pt(0)
            if i == 0:
                q.runs[0].bold = True
        self.doc.add_paragraph()
        if attention:
            self.doc.add_paragraph().add_run(f"Attention: {attention}").bold = True
        self.doc.add_paragraph("Dear Sir/Madam,")
        s = self.doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = s.add_run(subject.upper())
        run.bold, run.underline = True, True

    def para(self, text: str, bold: bool = False, italic: bool = False, size: float | None = None, align=None) -> None:
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def bullets(self, items: Iterable[str]) -> None:
        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

    def numbered(self, items: Iterable[str]) -> None:
        # Literal numbers: Word's shared "List Number" numbering would continue
        # across separate lists in the same document.
        for i, item in enumerate(items, start=1):
            q = self.doc.add_paragraph(f"{i}.	{item}")
            q.paragraph_format.left_indent = Cm(0.9)
            q.paragraph_format.first_line_indent = Cm(-0.6)
            q.paragraph_format.space_after = Pt(2)

    def heading(self, text: str, level: int = 1) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12 if level == 1 else 11)
        r.font.color.rgb = NAVY

    def closing(self, signatories: Sequence[tuple[str, str]], enclosures: Sequence[str] = (), cc: Sequence[str] = ()) -> None:
        self.doc.add_paragraph("Yours faithfully,")
        self.doc.add_paragraph(f"For: {self.profile.get('institution_name', '')}").runs[0].bold = True
        self.signatures(signatories)
        if enclosures:
            self.doc.add_paragraph().add_run("Enclosures:").bold = True
            self.numbered(enclosures)
        if cc:
            self.doc.add_paragraph().add_run("cc:").bold = True
            for line in cc:
                q = self.doc.add_paragraph(line)
                q.paragraph_format.space_after = Pt(0)

    def signatures(self, signatories: Sequence[tuple[str, str]]) -> None:
        # One block per person: a CCO who is also the reporting officer signs once.
        merged: dict[str, tuple[str, list[str]]] = {}
        for name, title in signatories:
            key = name.split(",")[0].strip().lower()
            if not key:
                continue
            if key in merged:
                merged[key][1].append(title)
            else:
                merged[key] = (name.split(",")[0].strip(), [title])
        sigs = [(name, " / ".join(titles)) for name, titles in merged.values()]
        if not sigs:
            return
        table = self.doc.add_table(rows=4, cols=len(sigs))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (name, title) in enumerate(sigs):
            cells = [table.cell(r, i) for r in range(4)]
            cells[0].paragraphs[0].add_run("\n\n_____________________________")
            cells[1].paragraphs[0].add_run(name).bold = True
            cells[2].paragraphs[0].add_run(title)
            cells[3].paragraphs[0].add_run("Date: ____________________")
        self.doc.add_paragraph()

    # ── Tables ────────────────────────────────────────────────────────────────

    def table(
        self,
        headers: Sequence[str],
        rows: Sequence[Sequence[str]],
        widths_cm: Sequence[float] | None = None,
        numeric_cols: Sequence[int] = (),
        bold_last: bool = False,
        font_size: float = 9,
        bold_rows: Iterable[int] = (),
    ) -> None:
        bold_set = set(bold_rows)
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold, run.font.size = True, Pt(font_size)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, "1F3A5F")
        # Repeat the header row on each page.
        tr_pr = t.rows[0]._tr.get_or_add_trPr()
        rh = OxmlElement("w:tblHeader")
        rh.set(qn("w:val"), "true")
        tr_pr.append(rh)
        for r_i, row in enumerate(rows):
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                run = p.add_run(str(value))
                run.font.size = Pt(font_size)
                if (bold_last and r_i == len(rows) - 1) or r_i in bold_set:
                    run.bold = True
                if i in numeric_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()

    def key_values(self, pairs: Sequence[tuple[str, str]]) -> None:
        t = self.doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in pairs:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = "", ""
            r = cells[0].paragraphs[0].add_run(k)
            r.bold, r.font.size = True, Pt(9)
            _shade(cells[0], "EEF2F7")
            cells[1].paragraphs[0].add_run(v or "—").font.size = Pt(9)
            cells[0].width, cells[1].width = Cm(6), Cm(10.5)
        self.doc.add_paragraph()

    def page_break(self) -> None:
        self.doc.add_page_break()

    def to_bytes(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


def abbreviation(name: str) -> str:
    stop = {"limited", "ltd", "plc", "microfinance", "bank", "mfb", "of", "the", "and", "&"}
    words = [w for w in "".join(c if c.isalnum() or c.isspace() else " " for c in name).split() if w.lower() not in stop]
    abbr = "".join(w[0] for w in words[:4]).upper()
    return (abbr or "MFB") + "MFB"
