"""Validate an operator's saved extraction before accepting durable evidence.

No OCR/model calls. Rich layouts stay in the immutable JSON artifact; canonical
page text/quality and exact-substring chunks use the existing database schema.
"""

import json
from copy import deepcopy

from ingestion.extraction import normalize, quality


def validate_cached(bundle, digest, file_type, path):
    if not isinstance(bundle, dict) or bundle.get("source", {}).get("sha256") != digest:
        raise ValueError("Cached extraction does not match the original checksum")
    serialized = json.dumps(bundle, ensure_ascii=False, allow_nan=False).encode("utf-8")
    if len(serialized) > 32 * 1024 * 1024:
        raise ValueError("Cached extraction exceeds the supported artifact size")
    records = bundle.get("pages")
    if not isinstance(records, list) or not records:
        raise ValueError("Cached extraction needs nonempty page/sheet records")
    if file_type == "pdf":
        import pypdfium2 as pdfium

        with pdfium.PdfDocument(path) as pdf:
            expected = len(pdf)
        if [item.get("page_number") for item in records] != list(range(1, expected + 1)):
            raise ValueError("Cached extraction must preserve every physical PDF page in order")
    elif file_type == "xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if [item.get("locator") for item in records] != workbook.sheetnames:
                raise ValueError("Cached extraction must preserve every spreadsheet sheet")
        finally:
            workbook.close()
        if any(item.get("page_number") is not None for item in records):
            raise ValueError("Spreadsheet sheets do not have physical PDF page numbers")
    else:
        raise ValueError("Cached extraction supports validated PDF and XLSX originals only")
    pages = []
    for item in records:
        raw, text = item.get("raw_text"), item.get("text")
        if not isinstance(raw, str) or not isinstance(text, str) or text != normalize(raw) or "\x00" in raw:
            raise ValueError("Cached text must match the documented normalization of raw text")
        method = item.get("method")
        if method not in {"pdf_text", "azure_layout", "windows_ocr_en_us", "openpyxl_cells"}:
            raise ValueError("Unsupported cached extraction method")
        metrics = deepcopy(item.get("quality", {}))
        recalculated = quality(raw)
        if any(metrics.get(key) != value for key, value in recalculated.items()):
            raise ValueError("Cached extraction quality counts do not match the saved text")
        if method == "windows_ocr_en_us":
            metrics["local_ocr_manual_review_required"] = True
            metrics["ocr_confidence_available"] = False
        confidence = metrics.get("ocr_confidence")
        if confidence is not None and (not isinstance(confidence, (int, float)) or not 0 <= confidence <= 1):
            raise ValueError("OCR confidence must be between zero and one")
        if method == "azure_layout" and confidence is None and text.strip():
            metrics["ocr_confidence_missing"] = True
        # Keep the full structured record in the artifact, not arbitrary columns.
        metrics["table_count"] = len(item.get("tables", []))
        pages.append({
            "page_number": item.get("page_number"), "locator": item.get("locator"),
            "method": method, "raw_text": raw, "text": text, "quality": metrics,
        })
    return pages, serialized
